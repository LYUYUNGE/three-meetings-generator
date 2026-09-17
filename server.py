from __future__ import annotations

import json
import base64
import html
import mimetypes
import re
import shutil
import socket
import subprocess
import sys
import tempfile
import threading
import webbrowser
from copy import deepcopy
from datetime import datetime
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from io import BytesIO
from pathlib import Path
from urllib.parse import quote, unquote, urlparse
from urllib.request import Request, urlopen

try:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt
    from docx.text.paragraph import Paragraph
except ImportError:
    print("缺少 python-docx。请使用 start.bat 启动本系统。", file=sys.stderr)
    raise


FROZEN = bool(getattr(sys,"frozen",False))
ROOT = Path(getattr(sys,"_MEIPASS",Path(__file__).resolve().parent))
WEB = ROOT / "web"
TEMPLATES = ROOT / "templates" / "board"
SHAREHOLDER_TEMPLATES = ROOT / "templates" / "shareholder"
if FROZEN:
    USER_ROOT=Path.home()/"AppData"/"Local"/"雷石股份"/"三会文件生成系统数据"
else:
    USER_ROOT=Path(__file__).resolve().parent
OUTPUTS = USER_ROOT / "outputs"
DATA = USER_ROOT / "data"
ARCHIVE = USER_ROOT / "archive"
HOST = "127.0.0.1"
PORT = 8765
MAX_BODY = 40 * 1024 * 1024

BOARD_BASE=[
 {"name":"王川","title":"董事长","identities":["董事"],"special_roles":["实际控制人","控股股东","董事长"]},
 {"name":"马杰","title":"董事、总经理","identities":["董事","高级管理人员"],"special_roles":["实际控制人","总经理"]},
 {"name":"张浩","title":"董事、副总经理","identities":["董事","高级管理人员"],"special_roles":["副总经理"]},
 {"name":"宗旭","title":"董事、副总经理","identities":["董事","高级管理人员"],"special_roles":["副总经理"]},
 {"name":"程鹏","title":"董事","identities":["董事"]},
 {"name":"孙承斌","title":"董事","identities":["董事"]},
 {"name":"王珂","title":"董事、副总经理","identities":["董事","高级管理人员"],"special_roles":["副总经理"]},
 {"name":"武剑","title":"副总经理","identities":["高级管理人员"],"special_roles":["副总经理"]},
 {"name":"李婧超","title":"财务负责人、董事会秘书","identities":["高级管理人员"],"special_roles":["董事会秘书","财务负责人"]},
]
SUPERVISORY_BASE=[
 {"name":"庄义峰","title":"监事会主席","identities":["监事"],"special_roles":["监事会主席"]},
 {"name":"刘三德","title":"监事","identities":["监事"]},
 {"name":"王晓伟","title":"监事","identities":["监事"]},
]
SHAREHOLDER_BASE=[
 {"name":"王川","shares":14738898},{"name":"马杰","shares":10131579},{"name":"邵志军","shares":1755065},
 {"name":"张浩","shares":1370291},{"name":"北京雷杰科技合伙企业（有限合伙）","shares":1157728},
 {"name":"石河子市尚势成长股权投资合伙企业（有限合伙）","shares":1063382},
 {"name":"银杏自清（天津）创业投资合伙企业（有限合伙）","shares":1033000},
 {"name":"北京雷岳科技合伙企业（有限合伙）","shares":591076},{"name":"西藏福茂实业有限公司","shares":529926},
 {"name":"宁波银盈成长股权投资合伙企业（有限合伙）","shares":455735},{"name":"陈彤","shares":303824},
 {"name":"于泓","shares":245928},{"name":"薛滨","shares":245928},{"name":"北京雷珏晨曦科技合伙企业（有限合伙）","shares":244558},
 {"name":"正和共创海河（天津）投资合伙企业（有限合伙）","shares":197485},{"name":"赵晋梅","shares":163255},
 {"name":"赖晓凌","shares":151912},{"name":"北京悦米科技有限公司","shares":126646},{"name":"宋欣儒","shares":100049},
 {"name":"李迎","shares":48984},{"name":"查倩","shares":47517},{"name":"蒋国庆","shares":35328},
 {"name":"蔡英武","shares":35328},{"name":"上海晨熹创业投资中心（有限合伙）","shares":31662},
 {"name":"北京君利联合创业投资合伙企业（有限合伙）","shares":30382},{"name":"崔绍辉","shares":29660},
 {"name":"李婧超","shares":24587},{"name":"南京银湖策略投资合伙企业（有限合伙）","shares":17841},
 {"name":"李石平","shares":17664},{"name":"姚刚","shares":17664},{"name":"宋一强","shares":17664},
 {"name":"张建华","shares":17664},{"name":"北京雷乾易墅科技合伙企业（有限合伙）","shares":16049},
 {"name":"汪彦","shares":14131},{"name":"冯义茂","shares":8832},
]

def shareholder_personnel():
    merged={x["name"]:{"name":x["name"],"title":"股东","identities":["股东"],"shares":x["shares"],"attending":True} for x in SHAREHOLDER_BASE}
    for person in BOARD_BASE+SUPERVISORY_BASE:
        item=merged.setdefault(person["name"],{"name":person["name"],"title":person.get("title",""),"identities":[],"shares":0,"attending":True})
        item["title"]=person.get("title",item.get("title",""))
        item["special_roles"]=list(person.get("special_roles",[]))
        for identity in person.get("identities",[]):
            if identity not in item["identities"]: item["identities"].append(identity)
    return list(merged.values())

def clean_reference_text(text, limit=18000):
    text = html.unescape(str(text or ""))
    text = re.sub(r"<script[\s\S]*?</script>|<style[\s\S]*?</style>", " ", text, flags=re.I)
    text = re.sub(r"<[^>]+>", "\n", text)
    text = re.sub(r"[\t\r ]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()[:limit]

def extract_reference_file(item):
    name = str(item.get("name", "参考资料"))
    raw = base64.b64decode(item.get("data", ""), validate=True)
    if len(raw) > 15 * 1024 * 1024: raise ValueError(f"{name}超过15MB，无法读取")
    suffix = Path(name).suffix.lower()
    if suffix in (".txt", ".csv", ".md", ".html", ".htm"):
        for enc in ("utf-8-sig", "gb18030", "utf-16"):
            try: return clean_reference_text(raw.decode(enc))
            except UnicodeDecodeError: pass
    if suffix == ".docx":
        doc = Document(BytesIO(raw))
        chunks = [p.text for p in doc.paragraphs]
        chunks += [" | ".join(c.text for c in row.cells) for table in doc.tables for row in table.rows]
        return clean_reference_text("\n".join(chunks))
    if suffix in (".xlsx", ".xlsm"):
        try:
            import openpyxl
            wb = openpyxl.load_workbook(BytesIO(raw), read_only=True, data_only=True)
            chunks=[]
            for ws in wb.worksheets:
                chunks.append(f"工作表：{ws.title}")
                for row in ws.iter_rows(values_only=True):
                    line=" | ".join(str(v) for v in row if v not in (None, ""))
                    if line: chunks.append(line)
                    if sum(map(len,chunks))>18000: break
            return clean_reference_text("\n".join(chunks))
        except Exception as exc: raise ValueError(f"无法读取Excel文件{name}：{exc}")
    if suffix == ".pdf":
        try:
            from pypdf import PdfReader
            return clean_reference_text("\n".join((p.extract_text() or "") for p in PdfReader(BytesIO(raw)).pages))
        except Exception as exc: raise ValueError(f"无法读取PDF文件{name}：{exc}")
    if suffix in (".png", ".jpg", ".jpeg", ".webp", ".bmp"):
        return f"图片参考资料：{name}（已作为参考附件留存，供预览及人工核对；未提取图片文字）"
    if suffix in (".doc", ".xls"):
        raise ValueError(f"{name}是旧版Office格式，请另存为docx或xlsx后上传")
    raise ValueError(f"暂不支持读取{name}的内容")

def fetch_reference_url(url):
    parsed=urlparse(url)
    if parsed.scheme not in ("http","https") or not parsed.hostname: raise ValueError("参考网址必须是有效的http或https地址")
    if parsed.hostname in ("localhost","127.0.0.1","::1") or parsed.hostname.endswith(".local"): raise ValueError("参考网址不能指向本机或内网地址")
    req=Request(url,headers={"User-Agent":"Mozilla/5.0 BoardDocsLocal/1.0"})
    with urlopen(req,timeout=12) as resp:
        raw=resp.read(2*1024*1024+1)
        if len(raw)>2*1024*1024: raise ValueError("参考网页超过2MB")
        charset=resp.headers.get_content_charset() or "utf-8"
    try: page=raw.decode(charset,errors="replace")
    except LookupError: page=raw.decode("utf-8",errors="replace")
    return clean_reference_text(page)

def save_uploaded_files(items, target, allowed_suffixes):
    target.mkdir(parents=True,exist_ok=True); saved=[]
    for item in items:
        name=Path(str(item.get("name",""))).name
        suffix=Path(name).suffix.lower()
        if not name or suffix not in allowed_suffixes: raise ValueError(f"不支持上传文件：{name or '未命名文件'}")
        raw=base64.b64decode(item.get("data",""),validate=True)
        if len(raw)>25*1024*1024: raise ValueError(f"{name}超过25MB")
        path=(target/safe_name(name)).resolve()
        if target.resolve() not in path.parents: raise ValueError("上传文件名不安全")
        path.write_bytes(raw);saved.append(path.name)
    return saved


DOCUMENT_SUFFIXES = ["会议议案", "会议通知", "通知回执", "签到表", "会议议程", "表决票", "会议记录", "会议决议", "计票情况"]
SIGNATURE_IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".webp", ".bmp"}
SIGNATURE_SCAN_SUFFIXES = SIGNATURE_IMAGE_SUFFIXES | {".pdf"}


def output_file_for_label(rec, label):
    folder = Path(rec.get("folder", ""))
    candidates = [x for x in rec.get("files", []) if label in str(x)]
    if not candidates: return None
    path = (folder / candidates[0]).resolve()
    try:
        if folder.resolve() not in path.parents or not path.is_file(): return None
    except Exception:
        return None
    return path


def append_signature_images_to_docx(source, image_paths, target):
    doc = Document(source)
    for path in image_paths:
        p = doc.add_paragraph()
        p.add_run().add_break(WD_BREAK.PAGE)
        pic = doc.add_paragraph()
        pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pic.add_run().add_picture(str(path), width=Cm(16))
    apply_document_fonts(doc)
    target.parent.mkdir(parents=True, exist_ok=True)
    doc.save(target)


def word_to_pdf(source, target):
    source = Path(source).resolve()
    target = Path(target).resolve()
    target.parent.mkdir(parents=True, exist_ok=True)
    script = r'''
param([string]$Source,[string]$Target)
$word = $null
$doc = $null
try {
  $word = New-Object -ComObject Word.Application
  $word.Visible = $false
  $doc = $word.Documents.Open($Source, $false, $true)
  $doc.ExportAsFixedFormat($Target, 17)
} finally {
  if ($doc -ne $null) { $doc.Close($false) }
  if ($word -ne $null) { $word.Quit() }
}
'''
    with tempfile.TemporaryDirectory() as tmp:
        script_path = Path(tmp) / "word_to_pdf.ps1"
        script_path.write_text(script, encoding="utf-8")
        cmd = [
            "powershell.exe",
            "-NoProfile",
            "-ExecutionPolicy",
            "Bypass",
            "-File",
            str(script_path),
            str(source),
            str(target),
        ]
        result = subprocess.run(cmd, capture_output=True, text=True, timeout=90)
    if result.returncode or not target.exists():
        detail = (result.stderr or result.stdout or "").strip()
        raise ValueError("无法将Word定稿版导出为PDF，请确认本机已安装Microsoft Word。" + (f" 详情：{detail}" if detail else ""))


def image_to_pdf(source, target):
    from PIL import Image
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas
    image = Image.open(source)
    if image.mode not in ("RGB", "RGBA"):
        image = image.convert("RGB")
    page_w, page_h = A4
    margin = 18
    img_w, img_h = image.size
    scale = min((page_w - margin * 2) / img_w, (page_h - margin * 2) / img_h)
    draw_w, draw_h = img_w * scale, img_h * scale
    x, y = (page_w - draw_w) / 2, (page_h - draw_h) / 2
    c = canvas.Canvas(str(target), pagesize=A4)
    c.drawImage(ImageReader(image), x, y, draw_w, draw_h, preserveAspectRatio=True, mask="auto")
    c.save()


def scans_to_pdf(scan_paths, target):
    from pypdf import PdfReader, PdfWriter
    writer = PdfWriter()
    with tempfile.TemporaryDirectory() as tmp:
        for index, scan in enumerate(scan_paths, 1):
            scan = Path(scan)
            pdf_path = scan
            if scan.suffix.lower() in SIGNATURE_IMAGE_SUFFIXES:
                pdf_path = Path(tmp) / f"scan-{index}.pdf"
                image_to_pdf(scan, pdf_path)
            reader = PdfReader(str(pdf_path))
            for page in reader.pages:
                writer.add_page(page)
        target.parent.mkdir(parents=True, exist_ok=True)
        with target.open("wb") as stream:
            writer.write(stream)


def signature_page_indexes(pdf_path):
    from pypdf import PdfReader
    reader = PdfReader(str(pdf_path))
    indexes = []
    keywords = ("签字页", "签署页", "签名", "签章", "投票人签字", "出席人签字", "与会人员签字", "董事签字", "监事签字")
    for index, page in enumerate(reader.pages):
        text = page.extract_text() or ""
        compact = re.sub(r"\s+", "", text)
        if any(keyword in compact for keyword in keywords):
            indexes.append(index)
    return indexes


def combine_body_with_signed_pages(body_pdf, signed_pages_pdf, target):
    from pypdf import PdfReader, PdfWriter
    body = PdfReader(str(body_pdf))
    signed = PdfReader(str(signed_pages_pdf))
    replace = set(signature_page_indexes(body_pdf))
    writer = PdfWriter()
    for index, page in enumerate(body.pages):
        if index not in replace:
            writer.add_page(page)
    for page in signed.pages:
        writer.add_page(page)
    target.parent.mkdir(parents=True, exist_ok=True)
    with target.open("wb") as stream:
        writer.write(stream)
    return len(replace), len(signed.pages)

def informative_reference_sentences(text, keywords, limit=4):
    candidates=re.split(r"(?<=[。；！？])|\n+", text)
    scored=[]
    for idx,s in enumerate(candidates):
        s=re.sub(r"\s+"," ",s).strip(" ：;；")
        if not 18<=len(s)<=220 or any(x in s for x in ("版权所有","网站地图","风险提示","免责声明")): continue
        score=sum(3 for k in keywords if k and k in s)+sum(1 for k in ("人民币","万元","亿元","日期","名称","比例","期限","地点","备案","批准") if k in s)
        scored.append((score,-idx,s))
    picked=[]
    for score,neg_idx,s in sorted(scored,reverse=True):
        if not any(s in item[2] or item[2] in s for item in picked): picked.append((score,neg_idx,s.rstrip("。")+"。"))
        if len(picked)>=limit: break
    # 采用参考材料原有的先后顺序，尽量保留范本的叙述逻辑。
    return [s for _,_,s in sorted(picked,key=lambda item:-item[1])]

def draft_proposal(rough, meeting_type, references):
    rough=re.sub(r"\s+"," ",rough).strip(" 。；;")
    if not rough: raise ValueError("请填写议案说明")
    organ={"board":"董事会","supervisory":"监事会","shareholder":"股东会"}.get(meeting_type,"董事会")
    source="\n".join(x for x in references if x).strip()
    year=(re.search(r"20\d{2}",rough+" "+source) or [str(datetime.now().year)])[0]
    combined=rough+" "+source[:12000]
    if re.search(r"半年|半年度报告",combined):
        title=f"关于公司《{year}年半年度报告》的议案"
        intro=f"根据《公司法》《证券法》以及全国中小企业股份转让系统有关业务规则等规定，公司结合{year}年上半年度实际经营情况，组织编制了《{year}年半年度报告》。"
        close=f"现将《{year}年半年度报告》提交{organ}审议，请各位审议。"
        keys=[year,"半年度","经营","财务","报告"]
    elif re.search(r"子公司|对外投资|设立|增资|投资",combined):
        place="香港" if "香港" in combined else "境外" if "境外" in combined else "相关地区"
        title=f"关于公司拟在{place}设立子公司并对外投资的议案" if "设立" in combined or "子公司" in combined else "关于公司对外投资的议案"
        intro="为落实公司经营发展规划，进一步完善业务布局并提升持续经营能力，公司拟实施本次对外投资事项。"
        close=f"本次投资事项尚需按照境内外有关规定办理备案、审批、登记或其他必要手续。现将上述事项提交{organ}审议，并提请授权公司管理层办理与本次投资相关的具体事宜。"
        keys=["投资","子公司",place,"人民币","万元","备案","批准"]
    elif re.search(r"任免|聘任|选举|辞任|辞职|换届",combined):
        title="关于公司人员任免的议案" if "换届" not in combined else f"关于公司{organ}换届选举的议案"
        intro="根据《公司法》《公司章程》及公司治理制度的有关规定，结合公司治理及经营管理需要，公司拟对相关人员安排作出调整。"
        close=f"相关人员的任职资格及任期按照法律法规、《公司章程》和公司内部制度执行。现将上述事项提交{organ}审议。"
        keys=["任免","聘任","选举","辞职","姓名","任期","职务"]
    elif re.search(r"关联交易|借款|担保|授信|贷款|财务资助",combined):
        subject="关联交易" if "关联交易" in combined else "担保" if "担保" in combined else "融资及借款"
        title=f"关于公司{subject}事项的议案"
        intro=f"为满足公司经营发展及资金安排需要，公司拟开展本次{subject}事项。公司将依据有关法律法规、《公司章程》及内部管理制度履行相应审议程序。"
        close=f"公司将根据实际情况签署相关协议并落实风险控制措施。现将上述事项提交{organ}审议。"
        keys=[subject,"金额","期限","利率","担保","关联方","用途"]
    else:
        core=re.split(r"[，。；;]",rough)[0]
        core=re.sub(r"^(公司|拟|审议通过|关于)","",core).strip()
        title=rough if rough.endswith("议案") and len(rough)<80 else f"关于{core}的议案"
        intro="根据相关法律法规、规范性文件及《公司章程》的有关规定，结合公司实际经营及治理需要，公司拟推进本议案所述事项。"
        close=f"现将上述事项提交{organ}审议，请各位审议。"
        keys=[x for x in re.split(r"\W+",rough) if len(x)>=2]
    details=informative_reference_sentences(source,keys)
    if details:
        # 参考资料存在时，直接沿用其事实展开顺序和段落式表达，不机械添加
        # “具体事项如下”或统一的编号清单。
        body=intro+"\n\n"+"\n\n".join(details)+"\n\n"+close
    else:
        body=intro+"\n\n"+rough.rstrip("。")+"。\n\n"+close
    if meeting_type in ("board","supervisory"):
        member="监事" if meeting_type=="supervisory" else "董事"
        body=re.sub(r"[，,]?\s*请各位(?:董事|监事)?审议[。.]?\s*$","",body).rstrip(" ，,。")+"。"
        body+=f"\n\n请各位{member}审议。"
    return {"title":title,"content":body,"reference_used":bool(source)}


def defaults_payload():
    items=records_load(); boards=[x for x in items if x.get("meeting_type")=="board" and not x.get("historical")]
    term,no=1,15
    roster=BOARD_BASE
    if boards:
        last=boards[-1]; inp=last.get("input",{}); term=int(inp.get("term_no",1)); no=int(inp.get("meeting_no",15))
        if last.get("status")=="signed_final":
            no+=1
            if any("换届" in (p.get("rough","")+p.get("title","")) for p in inp.get("proposals",[])): term,no=term+1,1
        for rec in reversed(boards):
            if rec.get("input",{}).get("personnel_change_confirmed"):
                roster=rec["input"].get("personnel_master") or rec["input"].get("participants") or roster; break
    roster=[x for x in deepcopy(roster) if x.get("name")!="王艳宇"]
    canonical={x["name"]:x for x in BOARD_BASE}
    for person in roster:
        current=canonical.get(person.get("name"))
        if current:
            person["title"]=current.get("title",person.get("title",""))
            person["identities"]=list(current.get("identities",[]))
            person["special_roles"]=list(current.get("special_roles",[]))
    supervisors=[x for x in items if x.get("meeting_type")=="supervisory" and not x.get("historical")]
    sterm,sno=1,1
    if supervisors:
        last=supervisors[-1]; inp=last.get("input",{}); sterm=int(inp.get("term_no",1)); sno=int(inp.get("meeting_no",1))
        if last.get("status")=="signed_final": sno+=1
    return {"board":{"term_no":term,"meeting_no":no,"personnel":roster},"supervisory":{"term_no":sterm,"meeting_no":sno,"personnel":SUPERVISORY_BASE},"shareholder":{"personnel":shareholder_personnel()},"stock_short":"雷石股份","stock_code":"875080"}

TEMPLATE_NAMES = {
    0: "0. 第一届董事会第十四次_会议议案.docx",
    1: "1. 第一届董事会第十四次_会议通知.docx",
    2: "2. 第一届董事会第十四次_通知回执.docx",
    3: "3. 第一届董事会第十四次_签到表.docx",
    4: "4. 第一届董事会第十四次_会议议程.docx",
    5: "5. 第一届董事会第十四次_表决票.docx",
    6: "6. 第一届董事会第十四次_会议记录.docx",
    7: "7. 第一届董事会第十四次_会议决议.docx",
}


def cn_date(value: str) -> str:
    value = (value or "").strip()
    if not value:
        return "____年__月__日"
    m = re.fullmatch(r"(\d{4})[-/](\d{1,2})[-/](\d{1,2})", value)
    if m:
        return f"{int(m.group(1))}年{int(m.group(2))}月{int(m.group(3))}日"
    return value


def safe_name(value: str) -> str:
    value = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", value).strip().rstrip(".")
    return value[:90] or "董事会文件"


def cn_num(n: int) -> str:
    digits = "零一二三四五六七八九"
    if n < 10:
        return digits[n]
    if n < 20:
        return "十" + (digits[n % 10] if n % 10 else "")
    if n < 100:
        return digits[n // 10] + "十" + (digits[n % 10] if n % 10 else "")
    return str(n)


def cn_digits(value) -> str:
    """Meeting ordinals use Chinese numerals while the form accepts Arabic digits."""
    try: return cn_num(int(value))
    except Exception: return str(value)


def set_header(doc, label: str, company: str, meeting_term: str) -> None:
    seen=set()
    for section in doc.sections:
        header=section.header
        key=id(header._element)
        if key in seen: continue
        seen.add(key)
        for table in list(header.tables):
            table._element.getparent().remove(table._element)
        for p in header.paragraphs: set_para(p,"")
        usable=section.page_width-section.left_margin-section.right_margin
        table=header.add_table(rows=1,cols=3,width=usable)
        trailing=header.paragraphs[0]
        trailing._element.addprevious(table._element)
        trailing.paragraph_format.space_before=Pt(0); trailing.paragraph_format.space_after=Pt(0); trailing.paragraph_format.line_spacing=Pt(1)
        for run in trailing.runs: run.font.size=Pt(1)
        table.autofit=False
        widths=[int(usable*0.46),int(usable*0.08),int(usable*0.46)]
        texts=[company,"",f"{meeting_term}{label}"]
        for i,(cell,width,text_value) in enumerate(zip(table.rows[0].cells,widths,texts)):
            cell.width=width
            tcpr=cell._tc.get_or_add_tcPr(); nowrap=OxmlElement("w:noWrap"); tcpr.append(nowrap)
            p=cell.paragraphs[0]; set_para(p,text_value); p.alignment=WD_ALIGN_PARAGRAPH.RIGHT if i==2 else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before=Pt(0); p.paragraph_format.space_after=Pt(0)
            for run in p.runs: run.font.size=Pt(8.5)
        tblpr=table._tbl.tblPr; borders=OxmlElement("w:tblBorders")
        for edge in ("top","left","bottom","right","insideH","insideV"):
            el=OxmlElement(f"w:{edge}")
            el.set(qn("w:val"),"nil")
            borders.append(el)
        tblpr.append(borders)


def apply_resolution_header_to_record(doc, company: str, meeting_term: str) -> None:
    reference=Document(TEMPLATES/TEMPLATE_NAMES[7])
    set_header(reference,"会议记录",company,meeting_term)
    add_header_bottom_line(reference)
    source=reference.sections[0].header._element
    seen=set()
    for section in doc.sections:
        header=section.header; key=id(header.part)
        if key in seen: continue
        seen.add(key)
        header.part._element=deepcopy(source)


def add_header_bottom_line(doc) -> None:
    seen=set()
    for section in doc.sections:
        header=section.header; key=id(header.part)
        if key in seen or not header.tables: continue
        seen.add(key)
        tblpr=header.tables[0]._tbl.tblPr
        borders=tblpr.find(qn("w:tblBorders"))
        if borders is None:
            borders=OxmlElement("w:tblBorders"); tblpr.append(borders)
        bottom=borders.find(qn("w:bottom"))
        if bottom is None:
            bottom=OxmlElement("w:bottom"); borders.append(bottom)
        bottom.set(qn("w:val"),"single"); bottom.set(qn("w:sz"),"4"); bottom.set(qn("w:space"),"0"); bottom.set(qn("w:color"),"808080")


def ensure_three_line_title(doc, company: str, meeting_term: str, label: str) -> None:
    if len(doc.paragraphs)<2: return
    set_para(doc.paragraphs[0],company); set_para(doc.paragraphs[1],meeting_term)
    normalized=lambda x: re.sub(r"\s+","",x)
    accepted={normalized(label),normalized(label.removeprefix("会议"))}
    if len(doc.paragraphs)>2 and normalized(doc.paragraphs[2].text) in accepted:
        set_para(doc.paragraphs[2],label)
    else:
        el=deepcopy(doc.paragraphs[1]._element); doc.paragraphs[1]._element.addnext(el); set_para(Paragraph(el,doc._body),label)
    reference=Document(TEMPLATES/TEMPLATE_NAMES[1])
    texts=(company,meeting_term,label)
    for i,text_value in enumerate(texts):
        target=doc.paragraphs[i]._element; styled=deepcopy(reference.paragraphs[i]._element)
        target.getparent().replace(target,styled); set_para(Paragraph(styled,doc._body),text_value)


def apply_document_fonts(doc) -> None:
    def style_run(run):
        run.font.name="Times New Roman"
        rpr=run._element.get_or_add_rPr(); fonts=rpr.get_or_add_rFonts()
        fonts.set(qn("w:ascii"),"Times New Roman"); fonts.set(qn("w:hAnsi"),"Times New Roman"); fonts.set(qn("w:eastAsia"),"宋体")
    for style in doc.styles:
        if hasattr(style,"font"):
            style.font.name="Times New Roman"
            rpr=style.element.get_or_add_rPr(); fonts=rpr.get_or_add_rFonts()
            fonts.set(qn("w:ascii"),"Times New Roman"); fonts.set(qn("w:hAnsi"),"Times New Roman"); fonts.set(qn("w:eastAsia"),"宋体")
    parts=[doc]
    seen=set()
    for section in doc.sections:
        for part in (section.header,section.footer):
            if id(part._element) not in seen: parts.append(part); seen.add(id(part._element))
    for part in parts:
        for p in part.paragraphs:
            for run in p.runs: style_run(run)
        for table in part.tables:
            for row in table.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        for run in p.runs: style_run(run)


def add_conditional_body_page_number(section) -> None:
    footer=section.footer; footer.is_linked_to_previous=False
    p=footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph(); set_para(p,""); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    def fld(kind, text=None):
        el=OxmlElement("w:fldChar" if kind!="instr" else "w:instrText")
        if kind=="instr": el.set(qn("xml:space"),"preserve"); el.text=text
        else: el.set(qn("w:fldCharType"),kind)
        return el
    r=p.add_run()._r
    for el in [fld("begin"),fld("instr"," IF "),fld("begin"),fld("instr"," SECTIONPAGES "),fld("separate"),fld("end"),fld("instr",' > 1 "'),fld("begin"),fld("instr"," PAGE "),fld("separate"),fld("end"),fld("instr",'" "" '),fld("separate"),fld("end")]: r.append(el)


def configure_body_only_page_numbers(doc) -> None:
    if len(doc.sections)<2: return
    for section in doc.sections[1:]:
        for ref in list(section._sectPr.findall(qn("w:footerReference"))): section._sectPr.remove(ref)
        section.footer.is_linked_to_previous=False
        for p in section.footer.paragraphs: set_para(p,"")
    add_conditional_body_page_number(doc.sections[0])


def set_para(p: Paragraph, text: str) -> None:
    if p.runs:
        p.runs[0].text = text
        for run in p.runs[1:]:
            run.text = ""
    else:
        p.add_run(text)


def set_real_paragraphs(p: Paragraph, text: str) -> list[Paragraph]:
    """Turn every visual line break into a real Word paragraph with cloned formatting."""
    parts=[x.strip() for x in re.split(r"\n+",str(text or "")) if x.strip()]
    if not parts:
        set_para(p,"")
        return [p]
    set_para(p,parts[0])
    result=[p]; cursor=p._element
    for part in parts[1:]:
        el=deepcopy(p._element)
        cursor.addnext(el)
        new_p=Paragraph(el,p._parent)
        set_para(new_p,part)
        result.append(new_p); cursor=el
    return result


def delete_para(p: Paragraph) -> None:
    el = p._element
    el.getparent().remove(el)


def replace_everywhere(doc, pairs, protected=()) -> None:
    protected_parts=[]
    for value in protected:
        protected_parts.extend(x for x in str(value).splitlines() if x)
    protected_parts=sorted(set(protected_parts),key=len,reverse=True)

    def converted(text):
        saved={}
        for i,value in enumerate(protected_parts):
            token=f"\ue000{i}\ue001"
            if value in text:
                text=text.replace(value,token);saved[token]=value
        for a,b in pairs: text=text.replace(a,b)
        for token,value in saved.items(): text=text.replace(token,value)
        return text

    parts = [doc]
    seen = set()
    for section in doc.sections:
        for part in (section.header, section.footer):
            key = id(part._element)
            if key not in seen:
                parts.append(part)
                seen.add(key)
    for part in parts:
        for p in part.paragraphs:
            old = p.text
            new = converted(old)
            if new != old:
                set_para(p, new)
        for table in part.tables:
            for row in table.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        old = p.text
                        new = converted(old)
                        if new != old:
                            set_para(p, new)


def find_para(doc, prefix: str) -> Paragraph:
    for p in doc.paragraphs:
        if p.text.startswith(prefix):
            return p
    raise ValueError(f"模板中未找到段落：{prefix}")


def replace_between(doc, start_prefix: str, end_prefix: str, lines: list[str]) -> None:
    ps = doc.paragraphs
    start = next(i for i, p in enumerate(ps) if p.text.startswith(start_prefix))
    end = next(i for i, p in enumerate(ps) if i > start and p.text.startswith(end_prefix))
    old = ps[start + 1 : end]
    if not old:
        raise ValueError(f"模板区间为空：{start_prefix}")
    prototype = deepcopy(old[0]._element)
    anchor = ps[end]._element
    for p in old:
        delete_para(p)
    for line in lines:
        el = deepcopy(prototype)
        anchor.addprevious(el)
        set_para(Paragraph(el, doc._body), line)


def rebuild_table_rows(table, rows: list[list[str]]) -> None:
    prototype = deepcopy(table.rows[1]._tr)
    while len(table.rows) > 1:
        table._tbl.remove(table.rows[-1]._tr)
    for values in rows:
        tr = deepcopy(prototype)
        table._tbl.append(tr)
        row = table.rows[-1]
        for i, value in enumerate(values):
            set_para(row.cells[i].paragraphs[0], str(value))


def common_pairs(data):
    return [
        ("北京雷石天地电子技术股份有限公司", data["company"]),
        ("第一届董事会第十四次", data["meeting_term"]),
        ("2026年8月14日", data["meeting_date_cn"]),
        ("2026年8月12日", data["notice_date_cn"]),
        ("王川", data["host"]),
        ("李婧超", data["contact_name"]),
        ("13426243126", data["contact_phone"]),
        ("北京市朝阳区富泽人寿大厦31层北京雷石天地电子技术股份有限公司", data["contact_address"]),
    ]


def confirmed_proposal_texts(data):
    """用户确认后的议案字段在任何会议类型中都不得参与模板替换。"""
    values=[]
    for proposal in data.get("proposals",[]):
        values.extend((proposal.get("title",""),proposal.get("content",""),proposal.get("result_note","")))
    return [value for value in values if value]


def lock_confirmed_proposal_fields(data):
    """用内部标记锁定用户确认文本；模板处理结束后再逐字恢复。"""
    locked=deepcopy(data);restore=[];serial=0
    for proposal in locked.get("proposals",[]):
        for field in ("title","content","result_note"):
            value=str(proposal.get(field,""))
            if not value: continue
            pieces=value.splitlines(keepends=True);rebuilt=[]
            for piece in pieces:
                ending="\n" if piece.endswith("\n") else ""
                text=piece[:-1] if ending else piece
                if text:
                    token=f"[[[CONFIRMED_{serial}]]]";serial+=1
                    restore.append((token,text));rebuilt.append(token+ending)
                else: rebuilt.append(ending)
            proposal[field]="".join(rebuilt)
    return locked,restore


def restore_confirmed_fields(paths, restore):
    if not restore: return
    for path in paths:
        doc=Document(path);replace_everywhere(doc,restore);doc.save(path)


def normalize(raw):
    required = ["company", "meeting_term", "notice_date", "meeting_date", "meeting_time", "location", "method"]
    for key in required:
        if not str(raw.get(key, "")).strip():
            raise ValueError(f"请填写必填字段：{key}")
    proposals = raw.get("proposals") or []
    if not proposals:
        raise ValueError("请至少填写一项议案")
    cleaned = []
    for i, p in enumerate(proposals, 1):
        title = str(p.get("title", "")).strip()
        content = str(p.get("content", "")).strip()
        if not title or not content:
            raise ValueError(f"第{i}项议案缺少名称或内容")
        cleaned.append({
            "title": title,
            "content": content,
            "shareholder": bool(p.get("shareholder")),
            "yes": "" if p.get("yes") in (None, "") else int(p.get("yes")),
            "no": "" if p.get("no") in (None, "") else int(p.get("no")),
            "abstain": "" if p.get("abstain") in (None, "") else int(p.get("abstain")),
            "recuse": "" if p.get("recuse") in (None, "") else int(p.get("recuse")),
            "result_note": str(p.get("result_note", p.get("record_after_vote", p.get("resolution_after_vote", "")))).strip(),
        })
    meeting_type = raw.get("meeting_type", "board")
    member_identity = "监事" if meeting_type == "supervisory" else "董事"
    member_default_title = "监事" if meeting_type == "supervisory" else "董事"
    participants=[x for x in raw.get("participants",[]) if str(x.get("name","")).strip() and x.get("attending",True)]
    directors = [x for x in raw.get("directors", []) if str(x.get("name", "")).strip()]
    if participants:
        directors=[{"name":x["name"],"title":x.get("title",member_default_title)} for x in participants if member_identity in x.get("identities",[])]
    if not directors:
        raise ValueError(f"请至少填写一名{member_identity}")
    for director in directors:
        if meeting_type == "supervisory":
            director["ballot_title"]="监事会主席" if "监事会主席" in str(director.get("title","")) else "监事"
        else:
            director["ballot_title"]="董事长" if "董事长" in str(director.get("title","")) else "董事"
    data = dict(raw)
    data["proposals"] = cleaned
    data["directors"] = directors
    data["executives"] = ([{"name":x["name"],"role":x.get("title","高级管理人员")} for x in participants if "高级管理人员" in x.get("identities",[])] if participants else [x for x in raw.get("executives", []) if str(x.get("name", "")).strip()])
    data["recorders"] = [x for x in raw.get("recorders", []) if str(x.get("name", "")).strip()]
    data["notice_date_cn"] = cn_date(raw["notice_date"])
    data["meeting_date_cn"] = cn_date(raw["meeting_date"])
    data["contact_name"] = str(raw.get("contact_name", "")).strip()
    data["contact_phone"] = str(raw.get("contact_phone", "")).strip()
    data["contact_address"] = str(raw.get("contact_address", "")).strip()
    data["convener"] = str(raw.get("convener", raw.get("host", ""))).strip()
    data["host"] = str(raw.get("host", "")).strip()
    default_host_role = "监事会主席" if meeting_type == "supervisory" else "董事长"
    data["host_role"] = str(raw.get("host_role", default_host_role)).strip() or default_host_role
    data["notice_delivery"] = str(raw.get("notice_delivery", "书面方式")).strip() or "书面方式"
    data["voting_methods"] = [str(x).strip() for x in raw.get("voting_methods",[]) if str(x).strip()] or ["现场投票"]
    data["expected_directors"] = int(raw.get("expected_directors") or len(directors))
    data["actual_directors"] = int(raw.get("actual_directors") or len(directors))
    data["duration"] = str(raw.get("duration", "半天")).strip() or "半天"
    if participants:
        roster_key = "supervisory_roster" if meeting_type == "supervisory" else "board_roster"
        all_board={x["name"] for x in raw.get(roster_key,[])}
        attending_board={x["name"] for x in participants if member_identity in x.get("identities",[])}
        all_exec={x["name"] for x in raw.get("executive_roster",[])}
        attending_exec=[x for x in participants if "高级管理人员" in x.get("identities",[])]
        bits=[]
        full_member_desc = "公司全体监事会成员" if meeting_type == "supervisory" else "公司全体董事会成员"
        partial_member_desc = "公司监事" if meeting_type == "supervisory" else "公司董事"
        bits.append(full_member_desc if all_board and attending_board==all_board else partial_member_desc+"、".join(x["name"] for x in participants if member_identity in x.get("identities",[])))
        if meeting_type != "supervisory" and all_exec and {x["name"] for x in attending_exec}==all_exec: bits.append("公司全体高级管理人员")
        elif meeting_type != "supervisory" and attending_exec: bits.extend(f"公司{x.get('title','高级管理人员')}{x['name']}" for x in attending_exec)
        others=[x for x in participants if not ({member_identity,"高级管理人员"}&set(x.get("identities",[])))]
        bits.extend(f"{x.get('title','其他')}{x['name']}" for x in others)
        data["participant_desc"]="、".join(bits)
    else: data["participant_desc"] = str(raw.get("participant_desc", "公司全体监事会成员" if meeting_type == "supervisory" else "公司全体董事会成员")).strip()
    return data


def build_proposals(data, out):
    doc = Document(TEMPLATES / TEMPLATE_NAMES[0])
    set_header(doc,"会议议案",data["company"],data["meeting_term"])
    replace_everywhere(doc, common_pairs(data), confirmed_proposal_texts(data))
    original = list(doc.paragraphs)
    block = [deepcopy(p._element) for p in original[:9]]
    body = doc._element.body
    sect_pr = body.sectPr
    for el in list(body):
        if el is not sect_pr:
            body.remove(el)
    for i, proposal in enumerate(data["proposals"], 1):
        elements = deepcopy(block if i < len(data["proposals"]) else block[:7])
        for el in elements:
            body.insert(len(body) - 1, el)
        paras = [Paragraph(el, doc._body) for el in elements]
        set_para(paras[0], f"议案{cn_num(i)}：{proposal['title']}")
        set_para(paras[1], "各位董事：")
        set_real_paragraphs(paras[2], proposal["content"])
        delete_para(paras[3])
        set_para(paras[5], data["company"])
        set_para(paras[6], "董事会")
    apply_document_fonts(doc); doc.save(out)


def build_notice(data, out):
    doc = Document(TEMPLATES / TEMPLATE_NAMES[1]); replace_everywhere(doc, common_pairs(data), confirmed_proposal_texts(data))
    set_header(doc,"会议通知",data["company"],data["meeting_term"])
    ensure_three_line_title(doc,data["company"],data["meeting_term"],"会议通知")
    set_para(find_para(doc, "时间："), f"时间：{data['meeting_date_cn']}{data['meeting_time']}")
    set_para(find_para(doc, "地点："), f"地点：{data['location']}")
    set_para(find_para(doc, "期限："), f"期限：{data['duration']}")
    vote_method="、".join(data.get("voting_methods") or ["现场投票"])
    set_para(find_para(doc, "二、会议表决方式"), f"二、会议召开及表决方式：本次会议采用{data['method']}召开，采用{vote_method}方式表决。")
    lines = [f"{i}、{p['title']}" for i, p in enumerate(data["proposals"], 1)]
    replace_between(doc, "三、拟审议的事项：", "四、会议召集人", lines)
    people = data["host"] if data["convener"] == data["host"] else f"召集人：{data['convener']}；主持人：{data['host']}"
    set_para(find_para(doc, "四、会议召集人"), f"四、会议召集人和主持人：{people}")
    set_para(find_para(doc, "联系人："), f"联系人：{data['contact_name']}")
    set_para(find_para(doc, "联系电话："), f"联系电话：{data['contact_phone']}")
    set_para(find_para(doc, "地址："), f"地址：{data['contact_address']}")
    intro = find_para(doc, "本人因")
    set_para(intro, f"本人因____________________原因无法参加{data['company']}（以下简称“公司”）{data['meeting_term']}会议，兹委托________先生/女士（居民身份证号码：_________________）代表本人参加本次会议，并授权其对公司于{data['notice_date_cn']}发出的《{data['company']}{data['meeting_term']}会议通知》中所列议案表决如下：")
    ps = doc.paragraphs
    start = next(i for i,p in enumerate(ps) if p._element is intro._element)
    end = next(i for i,p in enumerate(ps) if i > start and p.text.startswith("上述议案的简要意见"))
    old = ps[start+1:end]
    proto = deepcopy(old[0]._element)
    anchor = ps[end]._element
    for p in old: delete_para(p)
    for i,p in enumerate(data["proposals"],1):
        el=deepcopy(proto); anchor.addprevious(el)
        set_para(Paragraph(el,doc._body),f"{i}、{p['title']}：投（赞成/反对/弃权/回避）票")
    apply_document_fonts(doc); doc.save(out)


def build_receipt(data, out):
    doc=Document(TEMPLATES/TEMPLATE_NAMES[2]); set_header(doc,"通知回执",data["company"],data["meeting_term"]); replace_everywhere(doc,common_pairs(data),confirmed_proposal_texts(data)); ensure_three_line_title(doc,data["company"],data["meeting_term"],"通知回执"); apply_document_fonts(doc); doc.save(out)


def build_signin(data, out):
    doc=Document(TEMPLATES/TEMPLATE_NAMES[3]); set_header(doc,"签到表",data["company"],data["meeting_term"]); replace_everywhere(doc,common_pairs(data),confirmed_proposal_texts(data))
    rows=[[i,x["name"],x.get("title","董事"),""] for i,x in enumerate(data["directors"],1)]
    rebuild_table_rows(doc.tables[0],rows); ensure_three_line_title(doc,data["company"],data["meeting_term"],"签到表"); apply_document_fonts(doc); doc.save(out)


def build_agenda(data, out):
    doc=Document(TEMPLATES/TEMPLATE_NAMES[4]); set_header(doc,"会议议程",data["company"],data["meeting_term"]); replace_everywhere(doc,common_pairs(data),confirmed_proposal_texts(data))
    set_para(find_para(doc,"会议时间："),f"会议时间：{data['meeting_date_cn']}{data['meeting_time']}")
    set_para(find_para(doc,"会议地点："),f"会议地点：{data['location']}")
    set_para(find_para(doc,"参加人员："),f"参加人员：{data['participant_desc']}")
    set_para(find_para(doc,"主持人："),f"主持人：{data['host']}")
    lines=[f"{i}、{p['title']}" for i,p in enumerate(data["proposals"],1)]
    replace_between(doc,"一、议案宣读：","二、董事对议案",lines)
    mode_el=deepcopy(find_para(doc,"会议地点：")._element)
    anchor=find_para(doc,"参加人员：")._element; anchor.addprevious(mode_el)
    set_para(Paragraph(mode_el,doc._body),f"会议方式：{data['method']}")
    ensure_three_line_title(doc,data["company"],data["meeting_term"],"会议议程"); apply_document_fonts(doc); doc.save(out)


def build_ballot(data, out):
    doc=Document(TEMPLATES/TEMPLATE_NAMES[5]); set_header(doc,"表决票",data["company"],data["meeting_term"]); replace_everywhere(doc,common_pairs(data),confirmed_proposal_texts(data))
    rows=[[i,p["title"],"","","",""] for i,p in enumerate(data["proposals"],1)]
    rebuild_table_rows(doc.tables[0],rows)
    ensure_three_line_title(doc,data["company"],data["meeting_term"],"表决票")
    note=next((p for p in doc.paragraphs if p.text.strip().startswith("注：")),None)
    if note and re.match(r"^注：\s*1[、.]",note.text.strip()):
        rest=re.sub(r"^注：\s*","",note.text.strip()); set_para(note,"注：")
        el=deepcopy(note._element); note._element.addnext(el); first=Paragraph(el,doc._body); set_para(first,rest)
        second=next((p for p in doc.paragraphs if p.text.strip().startswith("2、")),None)
        for p in (note,first,second):
            if p:
                p.paragraph_format.left_indent=Cm(0)
                p.paragraph_format.first_line_indent=Cm(0.74)
    apply_document_fonts(doc); doc.save(out)


def vote_text(p):
    if all(p.get(k) in (None,"") for k in ("yes","no","abstain","recuse")):
        return "表决结果：赞成____票，反对____票，弃权____票，回避____票。\n\n"
    text=f"表决结果：赞成 {p['yes']} 票，反对 {p['no']} 票，弃权 {p['abstain']} 票"
    if p.get("recuse") not in (None,"") and int(p.get("recuse",0)):
        text+=f"，回避 {int(p['recuse'])} 票"
    return text


def rebuild_signature_pages(doc, start_index, block_len, signers, meeting_term, meeting_date, kind, prototype_elements=None):
    ps=list(doc.paragraphs)
    prototype=prototype_elements or [deepcopy(x._element) for x in ps[start_index:start_index+block_len]]
    body=doc._element.body; sect=body.sectPr
    for p in list(doc.paragraphs)[start_index:]: delete_para(p)
    while doc.paragraphs and not doc.paragraphs[-1].text.strip(): delete_para(doc.paragraphs[-1])
    for child in list(body):
        if child is sect: continue
        for old_sect in list(child.iter(qn("w:sectPr"))): old_sect.getparent().remove(old_sect)
    for el in prototype:
        for old_sect in list(el.iter(qn("w:sectPr"))): old_sect.getparent().remove(old_sect)
        for br in list(el.iter(qn("w:br"))):
            if br.get(qn("w:type"))=="page": br.getparent().remove(br)
    section_p=OxmlElement("w:p"); ppr=OxmlElement("w:pPr"); body_sect=deepcopy(sect)
    for old_type in list(body_sect.findall(qn("w:type"))): body_sect.remove(old_type)
    section_type=OxmlElement("w:type"); section_type.set(qn("w:val"),"nextPage"); body_sect.append(section_type)
    ppr.append(body_sect); section_p.append(ppr); body.insert(len(body)-1,section_p)
    for signer_no, signer in enumerate(signers):
        if signer_no:
            pb=OxmlElement("w:p"); run=OxmlElement("w:r"); br=OxmlElement("w:br"); br.set(qn("w:type"),"page"); run.append(br); pb.append(run); body.insert(len(body)-1,pb)
        els=deepcopy(prototype)
        for el in els: body.insert(len(body)-1,el)
        paras=[Paragraph(el,doc._body) for el in els]
        for p in paras:
            if p.text.startswith("（本页无正文"):
                set_para(p,f"（本页无正文，为{signer['company']}{meeting_term}会议{kind}签字页）")
            elif "签名：" in p.text:
                set_para(p,f"{signer['role']}签名：____________________")
            elif p.text.strip() in ({x["name"] for x in BOARD_BASE+SUPERVISORY_BASE}|{x["name"] for x in signers}):
                set_para(p,signer["name"])
            elif re.fullmatch(r"2026年\d+月\d+日",p.text.strip()):
                set_para(p,meeting_date)


def build_record(data, out):
    doc=Document(TEMPLATES/TEMPLATE_NAMES[6]); set_header(doc,"会议记录",data["company"],data["meeting_term"]); replace_everywhere(doc,common_pairs(data),confirmed_proposal_texts(data))
    ensure_three_line_title(doc,data["company"],data["meeting_term"],"会议记录")
    set_para(find_para(doc,"二、会议时间："),f"二、会议时间：{data['meeting_date_cn']}{data['meeting_time']}")
    set_para(find_para(doc,"三、会议地点："),f"三、会议地点及召开方式：{data['location']}，{data['method']}")
    set_para(find_para(doc,"四、参会人员："),f"四、参会人员：{data['participant_desc']}")
    set_para(find_para(doc,"五、会议主持人："),f"五、会议主持人：{data['host']}")
    lines=[]
    for i,p in enumerate(data["proposals"],1):
        lines.extend([f"{i}、审议：{p['title']}",vote_text(p)])
        lines.extend(x.strip() for x in re.split(r"\n+",p.get("result_note", "")) if x.strip())
    replace_between(doc,"（二）公司董事会成员","（三）主持人宣布",lines)
    signers=[]
    member_role="监事" if data.get("meeting_type")=="supervisory" else "董事"
    for x in data["directors"]: signers.append({"company":data["company"],"role":member_role,"name":x["name"]})
    if data.get("meeting_type")!="supervisory":
        for x in data["executives"]: signers.append({"company":data["company"],"role":x.get("role","高级管理人员"),"name":x["name"]})
        for x in data["recorders"]: signers.append({"company":data["company"],"role":"会议记录人","name":x["name"]})
    signature_start=next(i for i,p in enumerate(doc.paragraphs) if p.text.startswith("（本页无正文，为") and "会议记录签字页" in p.text)
    rebuild_signature_pages(doc,signature_start,12,signers,data["meeting_term"],data["meeting_date_cn"],"记录")
    apply_resolution_header_to_record(doc,data["company"],data["meeting_term"])
    configure_body_only_page_numbers(doc); apply_document_fonts(doc); doc.save(out)


def build_resolution(data, out):
    doc=Document(TEMPLATES/TEMPLATE_NAMES[7]); set_header(doc,"会议决议",data["company"],data["meeting_term"]); replace_everywhere(doc,common_pairs(data),confirmed_proposal_texts(data))
    ensure_three_line_title(doc,data["company"],data["meeting_term"],"会议决议")
    intro=doc.paragraphs[3]
    set_para(intro,f"{data['company']}（以下简称“公司”）于{data['meeting_date_cn']}以{data['method']}召开了{data['meeting_term']}会议，现场会议地点为{data['location']}。本次会议的通知已于{data['notice_date_cn']}以{data['notice_delivery']}送达全体董事。本次会议由{data['host']}{data['host_role']}主持，应到董事{data['expected_directors']}人，实到董事{data['actual_directors']}人。本次会议的召集和召开符合《公司法》和《公司章程》的有关规定，经审议并表决，所作决议合法有效。会议通过了如下决议：")
    lines=[]
    for i,p in enumerate(data["proposals"],1):
        lines.extend([f"{i}、审议：{p['title']}",vote_text(p)])
        lines.extend(x.strip() for x in re.split(r"\n+",p.get("result_note", "")) if x.strip())
    replace_between(doc,f"{data['company']}（以下简称","（以下无正文",lines)
    signers=[{"company":data["company"],"role":"董事","name":x["name"]} for x in data["directors"]]
    signature_start=next(i for i,p in enumerate(doc.paragraphs) if p.text.startswith("（本页无正文，为") and "会议决议签字页" in p.text)
    record_proto=Document(TEMPLATES/TEMPLATE_NAMES[6]); replace_everywhere(record_proto,common_pairs(data),confirmed_proposal_texts(data))
    rp=next(i for i,p in enumerate(record_proto.paragraphs) if p.text.startswith("（本页无正文，为") and "会议记录签字页" in p.text)
    proto=[deepcopy(x._element) for x in record_proto.paragraphs[rp:rp+12]]
    for el in proto:
        for sect in list(el.iter("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}sectPr")):
            sect.getparent().remove(sect)
    rebuild_signature_pages(doc,signature_start,10,signers,data["meeting_term"],data["meeting_date_cn"],"决议",proto)
    set_header(doc,"会议决议",data["company"],data["meeting_term"])
    configure_body_only_page_numbers(doc); apply_document_fonts(doc); doc.save(out)


def records_load():
    path = DATA / "records.json"
    if not path.exists(): return []
    try: return json.loads(path.read_text(encoding="utf-8"))
    except Exception: return []


def records_save(items):
    DATA.mkdir(parents=True, exist_ok=True)
    (DATA / "records.json").write_text(json.dumps(items, ensure_ascii=False, indent=2), encoding="utf-8")


def meeting_name(raw):
    kind=raw.get("meeting_type","board")
    if kind=="board": return f"第{cn_digits(raw.get('term_no',1))}届董事会第{cn_digits(raw.get('meeting_no',1))}次会议"
    if kind=="supervisory": return f"第{cn_digits(raw.get('term_no',1))}届监事会第{cn_digits(raw.get('meeting_no',1))}次会议"
    year=raw.get("meeting_year",datetime.now().year)
    return f"{year}年年度股东会会议" if raw.get("shareholder_kind")=="annual" else f"{year}年第{cn_digits(raw.get('meeting_no',1))}次临时股东会会议"


def vote_text_blank(p):
    vals=[p.get("yes"),p.get("no"),p.get("abstain")]
    if all(v in (None,"") for v in vals): return "表决结果：赞成____票，反对____票，弃权____票。"
    return vote_text({**p,"yes":p.get("yes",0),"no":p.get("no",0),"abstain":p.get("abstain",0),"recuse":p.get("recuse",0)})


def shareholder_vote_stats(raw, proposal):
    """Calculate a shareholder vote only from shareholders marked as attending."""
    attending={
        str(x.get("name","")):int(x.get("shares") or 0)
        for x in raw.get("personnel_master",[])
        if x.get("attending",True) and "股东" in x.get("identities",[]) and str(x.get("name","")).strip()
    }
    votes={str(x.get("name","")):str(x.get("result","yes")) for x in proposal.get("shareholder_votes",[])}
    totals={"yes":0,"no":0,"abstain":0,"recuse":0}
    for shareholder,shares in attending.items():
        result=votes.get(shareholder,"yes")
        if result not in totals: result="yes"
        totals[result]+=shares
    entitled=totals["yes"]+totals["no"]+totals["abstain"]
    totals["entitled"]=entitled
    totals["attending"]=attending
    return totals


def shareholder_result_text(stats):
    denom=stats["entitled"]
    def shares(value): return f"{int(value):,}股"
    def pct(value):
        if not denom: return "0%"
        value=f"{value*100/denom:.4f}".rstrip("0").rstrip(".")
        return value+"%"
    return (
        f"表决结果：赞成{shares(stats['yes'])}，占出席会议有表决权股份总数的{pct(stats['yes'])}；"
        f"反对{shares(stats['no'])}，占出席会议有表决权股份总数的{pct(stats['no'])}；"
        f"弃权{shares(stats['abstain'])}，占出席会议有表决权股份总数的{pct(stats['abstain'])}；"
        f"回避{shares(stats['recuse'])}，占出席会议有表决权股份总数的{pct(stats['recuse'])}。"
    )


def shareholder_attendees(raw):
    return [x for x in raw.get("personnel_master",[]) if x.get("attending",True)]


def replace_sequential(doc, predicate, values):
    targets=[p for p in doc.paragraphs if predicate(p.text.strip())]
    for p,value in zip(targets,values): set_para(p,value)


def shareholder_old_titles():
    doc=Document(sorted(SHAREHOLDER_TEMPLATES.glob("*.docx"))[0])
    return [p.text.split("：",1)[1] for p in doc.paragraphs if p.text.startswith("议案") and "：" in p.text]


def build_shareholder(raw, folder):
    name=meeting_name(raw); company=raw["company"]
    old="2026年第二次临时股东会会议"
    voting_methods=raw.get("voting_methods") or ["现场投票"]
    voting_text="、".join(voting_methods)+"方式"
    convener=raw.get("convener") or raw.get("convener_type") or "董事会"
    attendees=shareholder_attendees(raw)
    holders=[x for x in attendees if "股东" in x.get("identities",[])]
    represented=sum(int(x.get("shares") or 0) for x in holders)
    nonholders=[x for x in attendees if "股东" not in x.get("identities",[])]
    other_desc="、".join(f"{x.get('title','')}{x.get('name','')}" for x in nonholders) or "无"
    pairs=[("北京雷石天地电子技术股份有限公司",company),(old,name),("2026年第二次临时股东会",name.removesuffix("会议")),("2026年6月24日",cn_date(raw.get("notice_date"))),("2026年7月10日",cn_date(raw.get("meeting_date"))),("2026-7-10",raw.get("meeting_date","")),("10:00",raw.get("meeting_time","10:00")),("公司会议室",raw.get("location","公司会议室")),("现场会议",raw.get("method","现场会议")),("公司董事会",convener),("现场投票方式",voting_text),("现场投票",voting_text),("王川",raw.get("host","王川")),("董事长",raw.get("host_role","董事长")),("李婧超",raw.get("contact_name","李婧超")),("13426243126",raw.get("contact_phone","13426243126")),("北京市朝阳区富泽人寿大厦31层北京雷石天地电子技术股份有限公司",raw.get("contact_address","")),("本人因          原因","本人因____________________原因"),("雷石股份",raw.get("stock_short","雷石股份")),("875080",raw.get("stock_code","875080"))]
    props=raw.get("proposals") or []
    title_pairs=list(zip(shareholder_old_titles(),[p.get("title","") for p in props]))
    vote_results=[shareholder_vote_stats(raw,p) for p in props]
    result_lines=[shareholder_result_text(x) for x in vote_results]
    suffixes=["会议议案","会议通知","通知回执","签到表","会议议程","表决票","计票情况","会议记录","会议决议"]
    paths=[]
    srcs=sorted(SHAREHOLDER_TEMPLATES.glob("*.docx"))
    for i,(src,suffix) in enumerate(zip(srcs,suffixes)):
        doc=Document(src); set_header(doc,suffix,company,name); replace_everywhere(doc,pairs+title_pairs,confirmed_proposal_texts(raw))
        unused_titles=shareholder_old_titles()[len(props):]
        for paragraph in list(doc.paragraphs):
            if any(old_title in paragraph.text for old_title in unused_titles): delete_para(paragraph)
        # 将模板中的原议案名称替换为本次议案；多议案在议案正文中完整生成，其余文件先列明全部名称。
        old_titles=[p.text for p in doc.paragraphs if p.text.startswith("议案") and "：" in p.text]
        if old_titles and props:
            for j,t in enumerate(old_titles):
                if j<len(props): set_para(next(p for p in doc.paragraphs if p.text==t),f"议案{cn_num(j+1)}：{props[j]['title']}")
        if i==0 and props:
            # 保留模板排版，重建议案正文。
            original=list(doc.paragraphs); block=[deepcopy(p._element) for p in original[:9]]; body=doc._element.body; sect=body.sectPr
            for el in list(body):
                if el is not sect: body.remove(el)
            for j,pv in enumerate(props,1):
                els=deepcopy(block if j<len(props) else block[:7])
                for el in els: body.insert(len(body)-1,el)
                ps=[Paragraph(el,doc._body) for el in els]
                content=pv.get("content",pv.get("rough","")); set_para(ps[0],f"议案{cn_num(j)}：{pv['title']}"); set_para(ps[1],"各位股东："); set_real_paragraphs(ps[2],content)
                delete_para(ps[3])
                set_para(ps[5],company); set_para(ps[6],"董事会")
        if suffix=="会议通知":
            replace_sequential(doc,lambda t:t.startswith("（六）其他出席人员"),[f"（六）其他出席人员：{other_desc}"])
        if suffix in ("会议记录","会议决议"):
            result_paragraphs=[p for p in doc.paragraphs if p.text.strip().startswith(("表决结果：", "表决结果为：", "表决情况：", "经表决"))]
            for paragraph,result in zip(result_paragraphs,result_lines): set_para(paragraph,result)
            for paragraph in result_paragraphs[len(result_lines):]: delete_para(paragraph)
            for proposal,result in zip(props,result_lines):
                note=str(proposal.get("result_note","")).strip()
                if note:
                    target=next((p for p in doc.paragraphs if p.text==result),None)
                    if target: set_real_paragraphs(target,result+"\n"+note)
        if suffix=="计票情况" and doc.tables and props:
            table=doc.tables[0]
            rows=[]
            for index,(proposal,stats) in enumerate(zip(props,vote_results),1):
                rows.append([index,proposal.get("title",""),stats["yes"],stats["no"],stats["abstain"],stats["recuse"]])
            if len(table.rows)>1 and len(table.columns)>=6: rebuild_table_rows(table,rows)
        # Replace the old fixed attendance summary wherever it appears.
        for paragraph in doc.paragraphs:
            if "出席本次会议的股东/股东授权代表" in paragraph.text and "代表股份" in paragraph.text:
                set_para(paragraph,re.sub(r"出席本次会议的股东/股东授权代表.*?符合",f"出席本次会议的股东/股东授权代表共{len(holders)}名，代表股份{represented:,}股，符合",paragraph.text))
        if i>0: ensure_three_line_title(doc,company,name,suffix)
        apply_document_fonts(doc)
        out=folder/f"{i+1}. {safe_name(name)}_{suffix}.docx"; doc.save(out); paths.append(out)
    return paths


def build_meeting_email(raw):
    kind=raw.get("meeting_type","board")
    audience={"board":"董事","supervisory":"监事","shareholder":"股东"}.get(kind,"参会人员")
    office="监事会办公室" if kind=="supervisory" else "董事会办公室"
    short=raw.get("stock_short","雷石股份")
    term=raw.get("meeting_term") or meeting_name(raw)
    date=cn_date(raw.get("meeting_date"))
    time=str(raw.get("meeting_time","10:00"))
    try:
        hour,minute=(int(x) for x in time.split(":",1))
        period="上午" if hour<12 else "下午"
        display_hour=hour if hour<=12 else hour-12
        date_time=f"{date}{period}{display_hour}:{minute:02d}"
    except Exception:
        date_time=f"{date}{time}"
    methods=raw.get("meeting_methods") or []
    if methods: method_text="+".join(methods)
    else: method_text={"现场结合电子通讯方式":"现场会议+电子通讯会议","电子通讯方式":"电子通讯会议"}.get(raw.get("method"),raw.get("method",""))
    voting="、".join(raw.get("voting_methods") or ["现场投票"])
    proposals=raw.get("proposals") or []
    first_title=str(proposals[0].get("title","本次会议相关事项")) if proposals else "本次会议相关事项"
    topic=re.sub(r"^关于(?:公司)?|的议案$","",first_title).strip()
    if "报告" in topic:
        purpose=f"为全面履行公众公司的规范治理及信息披露义务，配合公司{topic}的正式报送"
    else:
        purpose=f"为做好{topic}的审议工作"
    proposal_lines="\n".join(f"{i}. {p.get('title','')}" for i,p in enumerate(proposals,1))
    return (
        f"尊敬的 {short} 各位{audience}：\n\n"
        f"您好！\n\n"
        f"{purpose}，公司拟定于{date_time}召开{term}。\n\n"
        f"现将本次会议的具体安排正式汇报如下：\n\n"
        f"【会议基本情况】\n\n"
        f"会议召开时间：{date} {time}\n"
        f"会议召开地点：{raw.get('location','')}\n"
        f"会议召开方式：{method_text}。\n"
        f"会议表决方式：{voting}\n\n"
        f"【拟审议议案】\n\n"
        f"本次会议拟审议以下{len(proposals)}项议案：\n\n{proposal_lines}\n\n"
        f"【附件及联系方式】\n\n"
        f"具体会议资料请您详阅本邮件附件。\n\n"
        f"感谢各位{audience}对公司各项工作的大力支持！\n\n"
        f"顺颂商祺\n\n"
        f"{raw.get('company','北京雷石天地电子技术股份有限公司')}\n"
        f"{office} {cn_date(raw.get('notice_date'))}"
    )


def generate(raw):
    raw=dict(raw); raw["meeting_term"]=meeting_name(raw)
    kind=raw.get("meeting_type","board")
    stamp=datetime.now().strftime("%Y%m%d-%H%M%S-%f")
    folder=OUTPUTS/f"{safe_name(raw['meeting_term'])}-{stamp}"; folder.mkdir(parents=True,exist_ok=False)
    if kind=="shareholder":
        locked_raw,restore=lock_confirmed_proposal_fields(raw)
        paths=build_shareholder(locked_raw,folder)
    else:
        data=normalize(raw)
        data,restore=lock_confirmed_proposal_fields(data)
        suffixes=["会议议案","会议通知","通知回执","签到表","会议议程","表决票","会议记录","会议决议"]
        paths=[folder/f"{i+1}. {safe_name(data['meeting_term'])}_{suffixes[i]}.docx" for i in range(8)]
        builders=[build_proposals,build_notice,build_receipt,build_signin,build_agenda,build_ballot,build_record,build_resolution]
        for fn,path in zip(builders,paths): fn(data,path)
        if kind=="supervisory":
            pairs=[("董事会","监事会"),("董事长","监事会主席"),("全体董事","全体监事"),("董事","监事")]
            protected=confirmed_proposal_texts(data)
            for path in paths:
                doc=Document(path); replace_everywhere(doc,pairs,protected); doc.save(path)
    restore_confirmed_fields(paths,restore)
    rid=stamp
    email=build_meeting_email(raw)
    rec={"id":rid,"created_at":datetime.now().isoformat(timespec="seconds"),"meeting_type":kind,"meeting_name":raw["meeting_term"],"status":"unclassified","input":raw,"folder":str(folder),"files":[p.name for p in paths],"email":email}
    items=records_load(); items.append(rec); records_save(items)
    return rec


class Handler(BaseHTTPRequestHandler):
    server_version="BoardDocsLocal/1.0"
    def log_message(self,fmt,*args): print("[本地]",fmt%args)
    def send_bytes(self,data,ctype,status=200,headers=None):
        self.send_response(status); self.send_header("Content-Type",ctype); self.send_header("Content-Length",str(len(data)))
        self.send_header("Cache-Control","no-store"); self.send_header("X-Content-Type-Options","nosniff")
        self.send_header("Content-Security-Policy","default-src 'self'; style-src 'self'; script-src 'self'; img-src 'self' data:; connect-src 'self'; object-src 'none'; frame-ancestors 'none'")
        for k,v in (headers or {}).items(): self.send_header(k,v)
        self.end_headers(); self.wfile.write(data)
    def do_GET(self):
        path=unquote(urlparse(self.path).path)
        if path=="/api/health": return self.send_bytes(b'{"ok":true,"offline":true}',"application/json; charset=utf-8")
        if path=="/api/records":
            return self.send_bytes(json.dumps(records_load(),ensure_ascii=False).encode("utf-8"),"application/json; charset=utf-8")
        if path=="/api/defaults":
            return self.send_bytes(json.dumps(defaults_payload(),ensure_ascii=False).encode("utf-8"),"application/json; charset=utf-8")
        if path.startswith("/api/files/"):
            bits=path.split("/",4)
            if len(bits)!=5: return self.send_bytes(b"Not found","text/plain",404)
            rid,filename=bits[3],bits[4]
            rec=next((x for x in records_load() if x["id"]==rid),None)
            if not rec: return self.send_bytes(b"Not found","text/plain",404)
            file=(Path(rec["folder"])/filename).resolve()
            if Path(rec["folder"]).resolve() not in file.parents or not file.is_file(): return self.send_bytes(b"Not found","text/plain",404)
            return self.send_bytes(file.read_bytes(),"application/vnd.openxmlformats-officedocument.wordprocessingml.document",200,{"Content-Disposition":f"attachment; filename*=UTF-8''{quote(file.name)}"})
        if path.startswith("/api/archive-files/"):
            bits=path.split("/",4)
            if len(bits)!=5: return self.send_bytes(b"Not found","text/plain",404)
            rid,filename=bits[3],bits[4]
            rec=next((x for x in records_load() if x["id"]==rid),None)
            if not rec: return self.send_bytes(b"Not found","text/plain",404)
            root=(Path(rec.get("archive_path",""))/"完整签署版PDF").resolve()
            file=(root/filename).resolve()
            if root not in file.parents or not file.is_file(): return self.send_bytes(b"Not found","text/plain",404)
            return self.send_bytes(file.read_bytes(),"application/pdf",200,{"Content-Disposition":f"attachment; filename*=UTF-8''{quote(file.name)}"})
        rel="index.html" if path=="/" else path.lstrip("/")
        file=(WEB/rel).resolve()
        if WEB.resolve() not in file.parents and file!=WEB.resolve(): return self.send_bytes(b"Forbidden","text/plain",403)
        if not file.is_file(): return self.send_bytes(b"Not found","text/plain",404)
        ctype=mimetypes.guess_type(file.name)[0] or "application/octet-stream"
        self.send_bytes(file.read_bytes(),ctype+("; charset=utf-8" if ctype.startswith("text/") or ctype=="application/javascript" else ""))
    def do_POST(self):
        try:
            length=int(self.headers.get("Content-Length","0"))
            if length<=0 or length>MAX_BODY: raise ValueError("提交内容大小不合法")
            raw=json.loads(self.rfile.read(length).decode("utf-8"))
            if self.path=="/api/preview-reference":
                if raw.get("file"):
                    preview=extract_reference_file(raw["file"])
                elif str(raw.get("url","")).strip():
                    preview=fetch_reference_url(str(raw["url"]).strip())
                else: raise ValueError("请选择参考资料或填写参考网址")
                return self.send_bytes(json.dumps({"preview":preview[:8000]},ensure_ascii=False).encode("utf-8"),"application/json; charset=utf-8")
            if self.path=="/api/draft":
                references=[]
                errors=[]
                for item in raw.get("files",[]):
                    try: references.append(extract_reference_file(item))
                    except Exception as exc: errors.append(str(exc))
                urls=raw.get("reference_urls") or [raw.get("reference_url","")]
                for url in dict.fromkeys(str(x).strip() for x in urls if str(x).strip()):
                    try: references.append(fetch_reference_url(url))
                    except Exception as exc: errors.append(f"参考网址读取失败（{url}）：{exc}")
                if errors: raise ValueError("；".join(errors))
                result=draft_proposal(str(raw.get("rough","")),str(raw.get("meeting_type","board")),references)
                return self.send_bytes(json.dumps(result,ensure_ascii=False).encode("utf-8"),"application/json; charset=utf-8")
            if self.path=="/api/delete-record":
                items=records_load(); rec=next((x for x in items if x.get("id")==raw.get("record_id")),None)
                if not rec: raise ValueError("未找到要删除的生成记录")
                allowed_root=(ARCHIVE if rec.get("historical") else OUTPUTS).resolve(); folder=Path(rec.get("folder","")).resolve()
                if folder.exists():
                    if allowed_root not in folder.parents: raise ValueError("记录对应的文件目录不安全，已停止删除")
                    shutil.rmtree(folder)
                archive_path=str(rec.get("archive_path","")).strip()
                if archive_path:
                    archive_root=ARCHIVE.resolve(); archived=Path(archive_path).resolve()
                    if archived.exists() and archived!=folder:
                        if archive_root not in archived.parents: raise ValueError("记录对应的归档目录不安全，已停止删除")
                        shutil.rmtree(archived)
                records_save([x for x in items if x.get("id")!=rec.get("id")])
                return self.send_bytes(json.dumps({"ok":True},ensure_ascii=False).encode("utf-8"),"application/json; charset=utf-8")
            if self.path=="/api/upload-signed-files":
                items=records_load();rec=next((x for x in items if x.get("id")==raw.get("record_id")),None)
                if not rec or rec.get("status")!="signed_final": raise ValueError("只有签署用定稿记录可以上传签署版本")
                root=Path(rec.get("archive_path") or (ARCHIVE/"签署定稿"/safe_name(rec["meeting_name"])/rec["id"])).resolve()
                if ARCHIVE.resolve() not in root.parents: raise ValueError("签署文件归档目录不安全")
                editable=save_uploaded_files(raw.get("editable",[]),root/"签署文件"/"可编辑版本",{".doc",".docx",".xls",".xlsx"})
                scanned=save_uploaded_files(raw.get("scanned",[]),root/"签署文件"/"定稿扫描件",{".pdf",".png",".jpg",".jpeg",".webp",".bmp"})
                signature_input=raw.get("signature_pages",{}) or {}
                signature_saved={}
                combined={}
                complete_pdfs={}
                complete_pdf_meta={}
                signature_root=root/"签字页扫描件"
                combined_root=root/"正文合并签字页"
                pdf_root=root/"完整签署版PDF"
                temp_pdf_root=root/"_临时PDF"
                for label,files in signature_input.items():
                    label=str(label)
                    if label not in DOCUMENT_SUFFIXES: continue
                    saved=save_uploaded_files(files,signature_root/safe_name(label),SIGNATURE_SCAN_SUFFIXES)
                    if saved: signature_saved[label]=saved
                    scan_paths=[signature_root/safe_name(label)/filename for filename in saved]
                    image_paths=[path for path in scan_paths if path.suffix.lower() in SIGNATURE_IMAGE_SUFFIXES]
                    source=output_file_for_label(rec,label)
                    if source and image_paths:
                        out=combined_root/f"{Path(source).stem}_正文合并签字页.docx"
                        append_signature_images_to_docx(source,image_paths,out)
                        combined[label]=out.name
                    if source and scan_paths:
                        body_pdf=temp_pdf_root/f"{Path(source).stem}_正文.pdf"
                        signed_pdf=temp_pdf_root/f"{safe_name(label)}_签字页.pdf"
                        final_pdf=pdf_root/f"{Path(source).stem}_完整签署版.pdf"
                        word_to_pdf(source,body_pdf)
                        scans_to_pdf(scan_paths,signed_pdf)
                        removed,inserted=combine_body_with_signed_pages(body_pdf,signed_pdf,final_pdf)
                        complete_pdfs[label]=final_pdf.name
                        complete_pdf_meta[label]={"replaced_unsigned_pages":removed,"inserted_signed_pages":inserted}
                if not editable and not scanned and not signature_saved: raise ValueError("请至少选择一份签署文件或签字页扫描件")
                existing_signed=rec.get("signed_uploads") or {"editable":[],"scanned":[]}
                existing_signed["editable"]=editable or existing_signed.get("editable",[])
                existing_signed["scanned"]=scanned or existing_signed.get("scanned",[])
                existing_signature=rec.get("signature_pages") or {}
                for label,names in signature_saved.items():
                    existing_signature[label]=names
                existing_combined=rec.get("combined_signed_files") or {}
                existing_combined.update(combined)
                existing_complete=rec.get("complete_signed_pdfs") or {}
                existing_complete.update(complete_pdfs)
                existing_complete_meta=rec.get("complete_signed_pdf_meta") or {}
                existing_complete_meta.update(complete_pdf_meta)
                rec["archive_path"]=str(root);rec["signed_uploads"]=existing_signed;rec["signature_pages"]=existing_signature;rec["combined_signed_files"]=existing_combined;rec["complete_signed_pdfs"]=existing_complete;rec["complete_signed_pdf_meta"]=existing_complete_meta;records_save(items)
                return self.send_bytes(json.dumps({"ok":True,"combined":combined,"complete_pdfs":complete_pdfs},ensure_ascii=False).encode("utf-8"),"application/json; charset=utf-8")
            if self.path=="/api/upload-history":
                kind=str(raw.get("meeting_type",""));term=int(raw.get("term_no") or 0);number=int(raw.get("meeting_no") or 0)
                if kind=="shareholder": raise ValueError("股东会历史会议目录暂未配置")
                if kind not in {"board","supervisory"} or term<1 or term>20 or number<1 or number>100: raise ValueError("历史会议届次不在开放范围内")
                editable=raw.get("editable",[]);scanned=raw.get("scanned",[])
                if not editable and not scanned: raise ValueError("请至少选择一份可编辑版本或签署版扫描件")
                organ="董事会" if kind=="board" else "监事会";name=f"第{cn_digits(term)}届{organ}第{cn_digits(number)}次会议";rid=f"historical-{kind}-{term}-{number}"
                target=(ARCHIVE/"历史会议"/safe_name(name)).resolve()
                if ARCHIVE.resolve() not in target.parents: raise ValueError("历史会议归档目录不安全")
                mode=str(raw.get("mode","append")); edit_dir=target/"可编辑版本"; scan_dir=target/"签署版扫描件"
                if mode=="replace":
                    for selected,path in ((editable,edit_dir),(scanned,scan_dir)):
                        if selected and path.exists(): shutil.rmtree(path)
                saved_editable=save_uploaded_files(editable,edit_dir,{".doc",".docx",".xls",".xlsx"})
                saved_scanned=save_uploaded_files(scanned,scan_dir,{".pdf",".png",".jpg",".jpeg",".webp",".bmp"})
                items=records_load();rec=next((x for x in items if x.get("historical") and x.get("meeting_type")==kind and int(x.get("input",{}).get("term_no",1))==term and int(x.get("input",{}).get("meeting_no",0))==number),None)
                uploads={"editable":[str(p) for p in sorted(edit_dir.glob("*")) if p.is_file()],"scanned":[str(p) for p in sorted(scan_dir.glob("*")) if p.is_file()]}
                if rec:
                    if mode=="replace":
                        for old in rec.pop("files",[]):
                            old_path=Path(old).resolve()
                            if old_path.exists() and target in old_path.parents: old_path.unlink()
                    rec["history_uploads"]=uploads;rec["created_at"]=datetime.now().isoformat(timespec="seconds")
                else:
                    rec={"id":rid,"created_at":datetime.now().isoformat(timespec="seconds"),"meeting_type":kind,"meeting_name":name,"status":"signed_final","historical":True,"input":{"meeting_type":kind,"term_no":term,"meeting_no":number,"proposals":[]},"folder":str(target),"archive_path":str(target),"history_uploads":uploads};items.append(rec)
                records_save(items)
                return self.send_bytes(json.dumps({"ok":True},ensure_ascii=False).encode("utf-8"),"application/json; charset=utf-8")
            if self.path=="/api/archive":
                items=records_load(); rec=next((x for x in items if x["id"]==raw.get("record_id")),None)
                if not rec: raise ValueError("未找到上次生成记录")
                status=raw.get("status")
                if status not in ("notice_final","signed_final","draft"): raise ValueError("请选择定稿状态")
                rec["status"]=status
                if status!="draft":
                    label="通知定稿" if status=="notice_final" else "签署定稿"
                    target=ARCHIVE/label/safe_name(rec["meeting_name"])/rec["id"]
                    target.parent.mkdir(parents=True,exist_ok=True)
                    if not target.exists(): shutil.copytree(rec["folder"],target)
                    rec["archive_path"]=str(target)
                    if status=="notice_final": rec["final_notice_date"]=raw.get("notice_date","")
                records_save(items)
                return self.send_bytes(json.dumps({"ok":True},ensure_ascii=False).encode(),"application/json; charset=utf-8")
            if self.path!="/api/generate": return self.send_bytes(b"Not found","text/plain",404)
            items=records_load()
            if items and items[-1].get("status")=="unclassified":
                return self.send_bytes(json.dumps({"ok":False,"needs_archive":True,"previous":items[-1]},ensure_ascii=False).encode("utf-8"),"application/json; charset=utf-8",409)
            rec=generate(raw)
            result={"ok":True,"record":rec,"files":[{"name":x,"url":f"/api/files/{rec['id']}/{quote(x)}"} for x in rec["files"]],"email":rec.get("email","")}
            self.send_bytes(json.dumps(result,ensure_ascii=False).encode("utf-8"),"application/json; charset=utf-8")
        except Exception as e:
            msg=json.dumps({"ok":False,"error":str(e)},ensure_ascii=False).encode("utf-8")
            self.send_bytes(msg,"application/json; charset=utf-8",400)


class ExclusiveThreadingHTTPServer(ThreadingHTTPServer):
    allow_reuse_address=False

    def server_bind(self):
        if hasattr(socket,"SO_EXCLUSIVEADDRUSE"):
            self.socket.setsockopt(socket.SOL_SOCKET,socket.SO_EXCLUSIVEADDRUSE,1)
        super().server_bind()


def main():
    missing=[name for name in TEMPLATE_NAMES.values() if not (TEMPLATES/name).exists()]
    if missing: raise SystemExit("模板缺失："+"、".join(missing))
    OUTPUTS.mkdir(parents=True,exist_ok=True); DATA.mkdir(parents=True,exist_ok=True); ARCHIVE.mkdir(parents=True,exist_ok=True)
    server=ExclusiveThreadingHTTPServer((HOST,PORT),Handler)
    print(f"公司三会文件生成系统已启动：http://{HOST}:{PORT}")
    print("仅允许本机访问；关闭此窗口即可停止。")
    threading.Timer(0.8,lambda:webbrowser.open(f"http://{HOST}:{PORT}")).start()
    try: server.serve_forever()
    except KeyboardInterrupt: pass
    finally: server.server_close()


if __name__=="__main__": main()
